"""Generate realistic SYNTHETIC term sheets for testing.

Real counterparty term sheets are confidential and not redistributable, so
these samples reproduce the structure, vocabulary and table layouts typical
of dealer term sheets for EQD structured products (FCN / BEN / Phoenix
autocall with KO & KI), with entirely fictitious trade details. Dealer
names are used only to exercise the masking agent.

Run:  python samples/generate_samples.py
"""
from pathlib import Path

from docx import Document as Docx
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)

OUT = Path(__file__).parent
styles = getSampleStyleSheet()
H = ParagraphStyle("H", parent=styles["Title"], fontSize=15, spaceAfter=4)
SUB = ParagraphStyle("SUB", parent=styles["Normal"], fontSize=9,
                     textColor=colors.grey, spaceAfter=10)
SMALL = ParagraphStyle("SMALL", parent=styles["Normal"], fontSize=7.5,
                       textColor=colors.grey, leading=9)

TABLE_STYLE = TableStyle([
    ("FONTSIZE", (0, 0), (-1, -1), 9),
    ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
    ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ("LINEBELOW", (0, 0), (-1, -2), 0.3, colors.HexColor("#cccccc")),
    ("TOPPADDING", (0, 0), (-1, -1), 3),
    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
])
UND_STYLE = TableStyle([
    ("FONTSIZE", (0, 0), (-1, -1), 8.5),
    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef1f5")),
    ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#bbbbbb")),
])

DISCLAIMER = ("This document is an indicative term sheet for discussion purposes only and does not "
              "constitute an offer to sell or a solicitation to buy any security. Structured products "
              "involve risk, including possible loss of principal. Capitalised terms are as defined in "
              "the final offering documentation. SYNTHETIC SAMPLE FOR SYSTEM TESTING — NOT A REAL TRADE.")


def pdf_term_sheet(filename, title, subtitle, rows, underlyings, contact_lines):
    doc = SimpleDocTemplate(str(OUT / filename), pagesize=A4,
                            topMargin=18 * mm, bottomMargin=15 * mm)
    story = [Paragraph(title, H), Paragraph(subtitle, SUB)]
    story.append(Table([[k, v] for k, v in rows], colWidths=[52 * mm, 118 * mm],
                       style=TABLE_STYLE))
    story.append(Spacer(1, 8))
    story.append(Paragraph("<b>Underlying Shares</b>", styles["Normal"]))
    story.append(Spacer(1, 3))
    und_data = [["Underlying", "Ticker", "Exchange", "Initial Price", "Strike Price"]] + underlyings
    story.append(Table(und_data, colWidths=[52 * mm, 28 * mm, 30 * mm, 30 * mm, 30 * mm],
                       style=UND_STYLE))
    story.append(Spacer(1, 10))
    for line in contact_lines:
        story.append(Paragraph(line, styles["Normal"]))
    story.append(Spacer(1, 8))
    story.append(Paragraph(DISCLAIMER, SMALL))
    doc.build(story)
    print("wrote", filename)


def docx_term_sheet(filename, title, subtitle, rows, underlyings, contact_lines):
    d = Docx()
    d.add_heading(title, level=1)
    d.add_paragraph(subtitle)
    t = d.add_table(rows=0, cols=2)
    t.style = "Light Grid Accent 1"
    for k, v in rows:
        cells = t.add_row().cells
        cells[0].text = k
        cells[1].text = v
    d.add_paragraph("")
    d.add_paragraph("Underlying Shares")
    u = d.add_table(rows=1, cols=5)
    u.style = "Light Grid Accent 1"
    hdr = u.rows[0].cells
    for i, h in enumerate(["Underlying", "Ticker", "Exchange", "Initial Price", "Strike Price"]):
        hdr[i].text = h
    for row in underlyings:
        cells = u.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    d.add_paragraph("")
    for line in contact_lines:
        d.add_paragraph(line)
    d.add_paragraph(DISCLAIMER)
    d.save(str(OUT / filename))
    print("wrote", filename)


# ------------------------------------------------------------ 1. Citi FCN
pdf_term_sheet(
    "FCN_Citi_WorstOf_3Stock_USD.pdf",
    "Fixed Coupon Note (FCN) — Worst-Of Basket",
    "Indicative Term Sheet | Private &amp; Confidential | Ref: EQD-FCN-2026-0187",
    [
        ("Issuer", "Citigroup Global Markets Holdings Inc."),
        ("Counterparty", "Citigroup Global Markets Singapore Pte. Ltd."),
        ("Product Type", "Fixed Coupon Note with Knock-Out and Knock-In (Autocallable)"),
        ("ISIN", "XS2911002876"),
        ("Trade Date", "06 July 2026"),
        ("Initial Fixing Date", "06 July 2026"),
        ("Issue Date", "13 July 2026"),
        ("Final Fixing Date", "06 July 2027"),
        ("Maturity Date", "13 July 2027"),
        ("Notional Amount", "USD 2,000,000"),
        ("Denomination", "USD 10,000"),
        ("Coupon Rate", "9.20% p.a., paid regardless of the performance of the Underlyings"),
        ("Coupon Frequency", "Monthly"),
        ("Strike Level", "80.00% of the Initial Price of each Underlying"),
        ("Knock-Out Barrier", "100.00% of the Initial Price"),
        ("Knock-Out Observation", "Monthly, on each Autocall Observation Date"),
        ("First Autocall Date", "06 October 2026"),
        ("Knock-In Barrier", "60.00% of the Initial Price"),
        ("Knock-In Observation", "Daily Close"),
        ("Basket Type", "Worst-Of"),
        ("Settlement Type", "Cash or Physical, at the election of the Issuer"),
        ("Settlement Currency", "USD"),
        ("Business Day Convention", "Modified Following"),
        ("Calculation Agent", "Citigroup Global Markets Limited"),
        ("Governing Law", "English Law"),
    ],
    [
        ["NVIDIA Corp", "NVDA UQ", "NASDAQ", "158.20", "126.56"],
        ["Advanced Micro Devices", "AMD UQ", "NASDAQ", "162.45", "129.96"],
        ["Broadcom Inc", "AVGO UQ", "NASDAQ", "271.30", "217.04"],
    ],
    ["Sales Contact: Amanda Lee", "Email: amanda.lee@citi-sales.example.com",
     "Phone: +65 6432 1188", "Settlement A/C: SG4488123456789012", "SWIFT: CITISGSGXXX"],
)

# ------------------------------------------------------------ 2. UBS BEN
pdf_term_sheet(
    "BEN_UBS_WorstOf_2Stock_SGD.pdf",
    "Bonus Enhanced Note (BEN)",
    "Indicative Terms and Conditions | UBS Structured Products | Ref: SPD-BEN-88412",
    [
        ("Issuer", "UBS AG, London Branch"),
        ("Counterparty", "UBS AG Singapore Branch"),
        ("Product Type", "Bonus Enhanced Note with Knock-In"),
        ("ISIN", "CH1300457221"),
        ("Trade Date", "08 July 2026"),
        ("Initial Fixing Date", "08 July 2026"),
        ("Issue Date", "15 July 2026"),
        ("Final Fixing Date", "08 January 2027"),
        ("Maturity Date", "15 January 2027"),
        ("Notional Amount", "SGD 1,500,000"),
        ("Denomination", "SGD 250,000"),
        ("Coupon Rate", "Bonus Coupon of 6.80% p.a."),
        ("Coupon Frequency", "At Maturity"),
        ("Strike Level", "100.00% of the Initial Price"),
        ("Knock-In Barrier", "75.00% of the Initial Price"),
        ("Knock-In Observation", "At Expiry"),
        ("Basket Type", "Worst-Of"),
        ("Settlement Type", "Cash"),
        ("Settlement Currency", "SGD"),
        ("Business Day Convention", "Following"),
        ("Calculation Agent", "UBS AG, London Branch"),
        ("Governing Law", "Swiss Law"),
    ],
    [
        ["DBS Group Holdings", "DBS SP", "SGX", "44.10", "44.10"],
        ["Singapore Telecommunications", "ST SP", "SGX", "3.42", "3.42"],
    ],
    ["Marketing Contact: Rajiv Menon", "Email: rajiv.menon@ubs-sales.example.com",
     "Phone: +65 6495 8000"],
)

# --------------------------------------------- 3. JPM Phoenix autocall
pdf_term_sheet(
    "Phoenix_JPMorgan_Autocall_HKD.pdf",
    "Phoenix Autocallable Note with Memory Coupon",
    "Indicative Term Sheet | J.P. Morgan Structured Investments | Ref: JPMSI-PHX-55107",
    [
        ("Issuer", "JPMorgan Chase Financial Company LLC"),
        ("Counterparty", "J.P. Morgan Securities (Asia Pacific) Limited"),
        ("Product Type", "Phoenix Autocallable with Knock-Out and Knock-In"),
        ("ISIN", "XS3055118840"),
        ("Trade Date", "01 July 2026"),
        ("Initial Fixing Date", "01 July 2026"),
        ("Issue Date", "08 July 2026"),
        ("Final Fixing Date", "01 July 2028"),
        ("Maturity Date", "10 July 2028"),
        ("Notional Amount", "HKD 8,000,000"),
        ("Denomination", "HKD 1,000,000"),
        ("Coupon Rate", "11.50% p.a., contingent, payable if the Worst Performing Underlying closes at or above the Coupon Barrier"),
        ("Coupon Frequency", "Quarterly"),
        ("Strike Level", "100.00% of the Initial Price"),
        ("Knock-Out Barrier", "100.00% of the Initial Price"),
        ("Knock-Out Observation", "Quarterly"),
        ("First Autocall Date", "01 October 2026"),
        ("Knock-In Barrier", "55.00% of the Initial Price"),
        ("Knock-In Observation", "Daily Close"),
        ("Basket Type", "Worst-Of"),
        ("Settlement Type", "Physical"),
        ("Settlement Currency", "HKD"),
        ("Business Day Convention", "Modified Following"),
        ("Calculation Agent", "J.P. Morgan Securities plc"),
        ("Governing Law", "English Law"),
    ],
    [
        ["Tencent Holdings", "700 HK", "HKEX", "412.60", "412.60"],
        ["Alibaba Group", "9988 HK", "HKEX", "108.30", "108.30"],
        ["Meituan", "3690 HK", "HKEX", "128.90", "128.90"],
    ],
    ["Prepared by: Kelvin Chan", "Email: kelvin.chan@jpm-desk.example.com",
     "Phone: +852 2800 1234", "Account: HK8821000045671122"],
)

# --------------------------------------------- 4. StanChart FCN (Word)
docx_term_sheet(
    "FCN_StandardChartered_Single_TSLA_USD.docx",
    "Fixed Coupon Note — Single Stock",
    "Indicative Term Sheet | Standard Chartered Bank | Ref: SCB-EQD-FCN-3391 | Private & Confidential",
    [
        ("Issuer", "Standard Chartered Bank (Singapore) Limited"),
        ("Counterparty", "Standard Chartered Bank"),
        ("Product Type", "Fixed Coupon Note with Knock-Out and Knock-In"),
        ("Trade Date", "09 July 2026"),
        ("Initial Fixing Date", "09 July 2026"),
        ("Issue Date", "16 July 2026"),
        ("Final Fixing Date", "09 April 2027"),
        ("Maturity Date", "16 April 2027"),
        ("Notional Amount", "USD 750,000"),
        ("Denomination", "USD 50,000"),
        ("Coupon Rate", "12.75% p.a."),
        ("Coupon Frequency", "Monthly"),
        ("Strike Level", "85.00% of the Initial Price"),
        ("Knock-Out Barrier", "103.00% of the Initial Price"),
        ("Knock-Out Observation", "Daily Close"),
        ("First Autocall Date", "09 August 2026"),
        ("Knock-In Barrier", "65.00% of the Initial Price"),
        ("Knock-In Observation", "Daily Close"),
        ("Basket Type", "Single"),
        ("Settlement Type", "Cash"),
        ("Settlement Currency", "USD"),
        ("Business Day Convention", "Modified Following"),
        ("Calculation Agent", "Standard Chartered Bank"),
        ("Governing Law", "English Law"),
    ],
    [["Tesla Inc", "TSLA UQ", "NASDAQ", "294.50", "250.33"]],
    ["Sales Contact: Priya Nair", "Email: priya.nair@scb-desk.example.com",
     "Phone: +65 6596 7000", "IBAN: GB29NWBK60161331926819"],
)

print("done")
